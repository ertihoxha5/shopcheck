from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out = Path('notes'); out.mkdir(exist_ok=True)
d = Document()
s=d.sections[0]; s.top_margin=s.bottom_margin=Inches(.7); s.left_margin=s.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st=d.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(11)
d.styles['Normal'].paragraph_format.space_after=Pt(7)
d.styles['Normal'].paragraph_format.line_spacing=1.08
d.styles['Title'].font.size=Pt(26)  
d.styles['Heading 1'].font.size=Pt(18)
d.styles['Heading 2'].font.size=Pt(13)
d.core_properties.title='ShopCheck shënime nga mësimet e para'
d.core_properties.author='ShopCheck learning notes prepared with AI assistance'
def p(t,style=None): return d.add_paragraph(t,style)
def h(t): d.add_heading(t,2)
def page(title): d.add_page_break(); d.add_heading(title,1)
def term(name,meaning,example,exercise):
    h(name); p(meaning); p('Shembull: '+example); p('Provo vetë: '+exercise)
def link(label,url):
    para=d.add_paragraph(); a=OxmlElement('w:hyperlink'); a.set(qn('r:id'),para.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
    r=OxmlElement('w:r'); props=OxmlElement('w:rPr'); c=OxmlElement('w:color'); c.set(qn('w:val'),'0563C1'); props.append(c); r.append(props); t=OxmlElement('w:t'); t.text=label; r.append(t); a.append(r); para._p.append(a)

d.add_heading('ShopCheck shënime nga mësimet e para',0)
p('Përgatitja dhe Phase 1', 'Subtitle')
p('Përmbledhje e bisedës deri më 1 tetor 2026. Këto janë shënime mësimore për t’i lexuar dhe rishkruar me fjalët e tua. Nuk janë raport i ekzekutimit të testeve.')
h('Ku kemi arritur')
p('Kemi kontrolluar mjetet bazë, kemi folur për kuptimin e QA dhe kemi lidhur Jira-n. Asistenti ka krijuar Epic-un e ShopCheck dhe një Task për kërkesat. Ende jemi në Phase 1; aplikacioni nuk është ndërtuar dhe nuk kemi ekzekutuar teste të tij.')
h('Qëllimi i projektit')
p('ShopCheck është një aplikacion i vogël për blerje online dhe projekti i parë i portofolit. Do të mësojmë si kalohet nga një kërkesë te testimi dhe te vendimi nëse versioni është gati. Dy projektet tjera ende nuk janë përcaktuar.')
p('Funksionet e planifikuara: login/logout, gjashtë produkte në dy kategori, kërkim dhe filtrim, shportë, kod zbritjeje, checkout me pagesë të simuluar dhe histori porosish.')
h('Si do të mësojmë')
p('Asistenti demonstron një shembull → një e bëjmë bashkë → një e provon ti vetë. Puna ndahet në hapa të vegjël. Jira menaxhohet nga asistenti sipas progresit real, ndërsa ti merr pjesë në shkrimin e kritereve dhe testimin manual.')
h('Rruga e planifikuar')
for t in ['1. Kërkesat dhe Jira — sqarojmë sjelljen dhe kriteret e pranimit.', '2. Aplikacioni — e ndërtojmë gradualisht.', '3. Dizajnimi i testeve — shkruajmë raste testimi.', '4. Ekzekutimi — testojmë, ruajmë dëshmi dhe verifikojmë rregullime.', '5. Automatizimi dhe release lokal — shtojmë pak teste automatike dhe raportin përfundimtar.']: p(t)

page('Mësimi 1 Kontrollimi i mjeteve')
p('Ambient i konfirmuar: Windows dhe PowerShell. Ti i ekzekutove komandat dhe dërgove screenshot; asistenti i lexoi rezultatet. Nuk u instalua asgjë në këtë hap.')
term('Node.js','Program që e ekzekuton JavaScript jashtë shfletuesit.','Në ShopCheck planifikojmë ta përdorim për serverin. Versioni i parë i verifikuar ishte v24.19.0.','Shkruaj me fjalët e tua për çka do të na duhet Node.js.')
term('Git','Mjet që ruan historinë e ndryshimeve të projektit.','Më vonë do ta lidhim një defekt dhe rregullimin e tij me versionin e kodit. Versioni i verifikuar ishte 2.42.0.windows.2.','Pse do të na ndihmonte ta dimë versionin ku u gjet një defekt?')
h('Komandat që i përdorëm')
p('Hape Start, shkruaj PowerShell dhe hape. Këto komanda mund të ekzekutohen në cilëndo dosje, një nga një:')
for t in ['node --version','git --version']:
    r=p(t).runs[0]; r.font.name='Consolas'
p('Rezultati i suksesshëm është numri i versionit. Komandat vetëm lexojnë informacion; nuk krijojnë skedarë dhe nuk instalojnë programe.')
p('Nëse shfaqet “is not recognized”, programi mund të mungojë ose PowerShell nuk po e gjen. Duhet lexuar mesazhi i saktë para se të zgjedhim zgjidhjen.')
h('Çka duhet ta kuptoj')
p('Të kesh Node.js dhe Git të instaluara nuk do të thotë që aplikacioni është ndërtuar ose testuar. Ky ishte vetëm kontroll i përgatitjes.')

page('Mësimi 2 Bazat e QA')
term('Quality Assurance ose QA','Puna që na ndihmon me parandalu dhe zbulu probleme, që produkti t’i plotësojë kërkesat. Nuk përbëhet vetëm prej klikimeve.','E sqarojmë rregullin e sasisë para se ta ndërtojmë shportën.','Përmend një problem të ShopCheck që mund ta dëmtojë blerësin.')
term('Test case','Udhëzim i përsëritshëm që tregon përgatitjen, veprimet dhe rezultatin që presim.','Për një shportë bosh, provojmë checkout dhe presim që porosia të mos krijohet. Ky është shembull i propozuar, jo test i ekzekutuar.','Cili është rezultati i pritshëm kur sasia e produktit është 0?')
term('Bug report','Raport për një defekt të konfirmuar, me hapa dhe dëshmi që dikush tjetër ta përsërisë.','Nëse checkout me shportë bosh krijon porosi, do ta hetonim kundrejt kërkesës. Ende nuk kemi zbuluar një defekt të tillë.','Pse fjalia “checkout nuk punon” nuk mjafton?')
h('Përgjigjja jote e parë')
p('Ti shkrove: “Për sasi 0 sistemi nuk duhet të lejoj checkout.” Arsyetimi ishte i drejtë: një porosi nuk duhet të pranojë sasi të pavlefshme.')
p('Asistenti propozoi sqarimin: ndryshimi i sasisë në 0 refuzohet; heqja e produktit është veprim më vete. Gjithashtu, shporta bosh nuk mund të përfundojë checkout. Ky sqarim ende duhet të futet në kërkesat dhe kriteret përkatëse.')
h('Rregulli i ndershmërisë')
p('Shembulli nuk është rezultat testimi. Dyshimi nuk është defekt i konfirmuar. Nuk shënojmë Pass ose Fail pa ekzekutim real dhe nuk i paraqesim kontrollet automatike të asistentit si testim manual të kryer nga ti.')

page('Mësimi 3 Organizimi në Jira')
term('Kanban board','Tabelë ku biletat organizohen në kolona sipas gjendjes së punës.','Board-i ynë është KAN board, në hapësirën shopcheck.','Hape board-in dhe gjeji kolonat që ekzistojnë.')
term('Epic','Grup i madh pune që bashkon funksione dhe detyra të lidhura.','KAN-4 përmbledh projektin ShopCheck.','Pse login dhe checkout mund të jenë nën të njëjtin Epic?')
term('Task','Biletë për një punë konkrete që duhet kryer.','KAN-5 është puna për sqarimin e kërkesave dhe rishikimin e kritereve.','Çka duhet të përfundojë para se KAN-5 të shënohet Done?')
term('Workflow','Rruga që ndjek puna nga fillimi deri në përfundim.','Propozimi ynë: To Do → In Progress → Ready for QA → In QA → Done.','A mjafton fjalia “e rregullova” për ta mbyllur një defekt?')
p('To Do = në pritje; In Progress = duke u punuar; Ready for QA = gati për testim; In QA = duke u testuar; Done = kriteret e përfundimit janë plotësuar.')
p('Ky është procesi ynë i portofolit, jo një proces i njohur i StarLabs. Board-i aktual ka To Do, Ready for QA, In QA dhe Done. In Progress mungon. Përkohësisht KAN-5 qëndron në To Do me etiketën in-progress dhe shpjegim në përshkrim.')
p('Llojet e konfirmuara në projekt janë Epic, Story, Task dhe Subtask. Bug nuk ishte në dispozicion gjatë kontrollit. Këto kufizime janë dokumentuar; nuk janë rregulluar ende. Story dhe Subtask janë parë në Jira, por ende nuk i kemi mësuar në hollësi.')

page('Mësimi 4 Kriteret e pranimit')
term('Acceptance criteria','Kushte të qarta dhe të kontrollueshme që tregojnë kur funksioni sillet siç kërkohet.','Pas rifreskimit, shporta e klientit ruan produktin dhe sasinë. Kjo sjellje është propozuar dhe ende po diskutohet.','Plotëso AC-2 më poshtë me një rezultat të saktë.')
h('Shembulli që e demonstroi asistenti')
p('AC-1: Given a signed-in customer has two units of a product in their cart, when they refresh the page, then the cart retains the same product and quantity.')
p('Given tregon gjendjen fillestare; When tregon veprimin; Then tregon rezultatin që presim. Këto tri pjesë na ndihmojnë me e bërë kriterin të testueshëm.')
h('Propozimi që ende pret diskutim')
p('Shporta ruhet veçmas për secilin klient pas rifreskimit dhe logout, dhe nuk u shfaqet klientëve të tjerë. Ky është propozim i asistentit, jo sjellje e implementuar ose kriter i aprovuar nga ti.')
h('Ushtrimi yt i papërfunduar')
p('AC-2: Given Customer A has products in their cart, when Customer A logs out and Customer B logs in, then …')
p('Përgjigjja ime në anglisht:')
p('______________________________________________________________________')
p('______________________________________________________________________')
p('Arsyetimi im në shqip:')
p('______________________________________________________________________')
p('______________________________________________________________________')
p('Pyetje ndihmëse: çka dimë për shportën që Customer B mund ta ketë pasur më herët? Formuloje kriterin pa lënë supozime të paqarta.')
p('Ky ushtrim ende nuk është dorëzuar ose vlerësuar. Pasi ta shkruash, do ta rishikojmë bashkë dhe asistenti do ta regjistrojë në Jira.')

page('Regjistri i punës dhe hapi tjetër')
h('Çka është bërë realisht')
for t in ['Ti kontrollove Node.js dhe Git në PowerShell dhe dërgove dëshmi me screenshot.', 'Ti ndave lidhjen dhe pamjen e Jira-s, aktivizove lidhjen dhe autorizove asistentin ta menaxhojë Jira-n për projektin.', 'Asistenti arriti ta lexojë Jira-n përmes Atlassian Rovo dhe verifikoi kolonat, llojet e biletave dhe biletat ekzistuese.', 'Asistenti krijoi KAN-4 si Epic dhe KAN-5 si Task të lidhur me të; krijimi u konfirmua nga Jira. Të dyja u verifikuan në To Do.', 'KAN-5 mori etiketën in-progress dhe përshkrimin e punës aktive. Nuk u mbyll asnjë biletë.']: p(t,'List Bullet')
link('KAN-4 — ShopCheck — Manual QA Portfolio','https://ertihoxha874.atlassian.net/browse/KAN-4')
link('KAN-5 — Define ShopCheck requirements and review acceptance criteria','https://ertihoxha874.atlassian.net/browse/KAN-5')
link('Hap board-in e ShopCheck','https://ertihoxha874.atlassian.net/jira/software/projects/KAN/boards/2')
p('Biletat paraprake u ruajtën pa ndryshim: KAN-1 “Task 1” është Task në To Do; KAN-2 “Task 2” është Story në Ready for QA; KAN-3 “Subtask 2.1” është Subtask në To Do. Qëllimi i tyre ende nuk është sqaruar. Statusi Ready for QA nuk është dëshmi që është kryer punë reale.')
h('Çka ende nuk është bërë')
p('Nuk kemi krijuar aplikacionin, qa/requirements.md, pesë stories të plota, test cases ose regjistra ekzekutimi. Nuk kemi kryer teste manuale të aplikacionit, raportuar defekte të konfirmuara, bërë rregullime apo reteste. Postman, Playwright dhe release janë hapa të ardhshëm, jo aftësi të ushtruara deri tash.')
h('Si t’i përdorësh këto shënime')
p('Lexoje një mësim dhe rishkruaje me fjalët e tua. Pastaj provo ushtrimin pa e kopjuar shpjegimin. Për hapin e radhës mjafton ta shkruash AC-2; nuk ke nevojë ta përfundosh krejt dokumentin menjëherë.')
p('Kontrollo veten: cili është dallimi mes një kriteri të pranimit dhe një rezultati testimi? Pse KAN-5 ende nuk është Done?')
p('Burimi: biseda jonë, screenshot-i yt i mjeteve dhe përgjigjet reale të Jira-s gjatë sesionit. Dokumenti u përgatit nga asistenti me AI për mësim; nuk dëshmon përvojë pune profesionale ose ekzekutime testesh.')

footer=s.footer.paragraphs[0]; footer.alignment=2
footer.add_run('ShopCheck • Shënime mësimore • ')
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
d.save(out/'ShopCheck_Lessons_01.docx')
print(out/'ShopCheck_Lessons_01.docx')
